#!/usr/bin/env python3
"""Refine PDF categorization to reach ~40 target."""

from openpyxl import load_workbook
import re

# Official/High-quality sources - ALWAYS KEEP
ALWAYS_KEEP = [
    "9789240003910",  # WHO
    "23-MH-8110",  # WHO
    "FS_Grief",  # WHO
    "MHF-UK",  # Mental Health Foundation UK
    "surgeon-general",  # US Surgeon General
    "FS_",  # Other WHO factsheets
]

# Priority keep patterns (clinical/research quality)
PRIORITY_KEEP = [
    r"DISTRESS TOLERANCE",
    r"EMOTION REGULATION",
    r"Problem-Management-Plus",
    r"Trauma-Informed",
    r"Preventing-suicide",
    r"clinical-research",
    r"Early-Adolescent",
    r"helping-children-and-adolescents-cope",
    r"Handbook-of-emotion",
    r"Fundamentals.*Psychological",
    r"Advanced.*anxiety",
    r"bipolar-disorder",
    r"obsessive-compulsive",
    r"post-traumatic-stress",
    r"schizophrenia",
    r"perinatal-depression",
    r"eating-disorders",
    r"psychosis",
    r"10\.54535",  # Research papers - select high-impact
    r"10\.18863",
    r"10\.26650",
    r"annurev-psych",  # Annual Review of Psychology
    r"turkpsikoloji",  # Turkish research
    r"resiliencepsikolojik",  # Turkish resilience
    r"Shortcuts-to-Social",  # Clinical manual
    r"Breakup-Book",  # Clinical guide
    r"Dealing.*Distress",  # Clinical guide
]

# Selective keep (good topic coverage)
SELECTIVE_KEEP = [
    r"self-esteem",
    r"relationships",
    r"attachment",
    r"loneliness",
    r"body-image",
    r"stress",
    r"grief",
    r"loss",
    r"rumination",
    r"rejection",
    r"boundaries",
    r"social-anxiety",
    r"emotion.*regulation",
    r"coping",
    r"distress.*tolerance",
]

# Delete patterns (low quality/redundant)
DELETE_PATTERNS = [
    r"^[a-f0-9]{8}-[a-f0-9]{4}",  # Hash filenames
    r"^[a-f0-9]{32,}",  # Long hashes
    r"^[0-9]{1,2}$",  # Just numbers
    r"IJIRT|qt0mb4f|ABUIABA9GAAgpLu",  # Spam-like filenames
]


def classify_pdf(filename: str, source_type: str, quality: str) -> str:
    """Classify a PDF as KEEP, DELETE, or UNCLEAR."""
    name_lower = filename.lower()

    # Official/highest quality - always keep
    if quality == "HIGH" or any(re.search(p, name_lower) for p in ALWAYS_KEEP):
        return "KEEP"

    # Priority sources - keep
    if any(re.search(p, name_lower) for p in PRIORITY_KEEP):
        return "KEEP"

    # Research papers - selective
    if source_type == "Research":
        # Keep high-impact or related to key topics
        if any(re.search(p, name_lower) for p in SELECTIVE_KEEP):
            return "KEEP"
        # Otherwise unclear - needs review
        return "UNCLEAR"

    # Self-help and clinical - generally keep if medium quality
    if source_type in ("SelfHelp", "Clinical") and quality == "MEDIUM":
        return "KEEP"

    # Turkish sources - keep
    if source_type == "Turkish":
        return "KEEP"

    # Low quality unknown - delete
    if quality == "LOW" or source_type == "Unknown":
        if any(re.search(p, name_lower) for p in DELETE_PATTERNS):
            return "DELETE"
        return "UNCLEAR"

    return "UNCLEAR"


def refine_audit():
    """Refine PDF classifications."""
    audit_path = "/Users/aykutakkus/Desktop/Projects/Capstone/data/pdf_audit.xlsx"
    wb = load_workbook(audit_path)
    ws = wb.active

    keep_count = 0
    delete_count = 0
    unclear_count = 0
    total_rows = 0

    # Skip summary rows at the end
    max_actual_row = ws.max_row
    for row in ws.iter_rows(min_row=ws.max_row - 5, max_row=ws.max_row):
        if row[0].value and "SUMMARY" in str(row[0].value):
            max_actual_row = row[0].row - 2
            break

    # Process each PDF
    for row in ws.iter_rows(min_row=2, max_row=max_actual_row):
        filename = row[0].value
        source_type = row[1].value
        quality = row[4].value

        if not filename or not isinstance(filename, str):
            continue

        total_rows += 1

        # Classify
        action = classify_pdf(filename, source_type or "Unknown", quality or "LOW")

        # Update action column (column F, index 5)
        row[5].value = action

        if action == "KEEP":
            keep_count += 1
        elif action == "DELETE":
            delete_count += 1
        else:
            unclear_count += 1

    # Save
    wb.save(audit_path)
    print(f"Refined categorization ({total_rows} PDFs):")
    print(f"  KEEP: {keep_count}")
    print(f"  DELETE: {delete_count}")
    print(f"  UNCLEAR: {unclear_count}")
    print(f"  Target: 40 PDFs")
    print(f"\nSaved: {audit_path}")


if __name__ == "__main__":
    refine_audit()
